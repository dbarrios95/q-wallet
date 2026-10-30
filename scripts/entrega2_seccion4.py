"""Genera las secciones 4.1, 4.2 y 4.5 del informe (Fase 2, bloque 1).

Uso:
    python scripts/entrega2_diagrama_flujo.py        # figura EVD-3-08
    python scripts/entrega2_seccion4.py               # Word de las secciones
    python scripts/entrega1_estilos.py docs/entregas/Entrega2_Seccion_4_Arquitectura_Controles.docx

Toma los estilos (títulos, márgenes, encabezado y número de página) del informe de la
Entrega 1 y vacía su contenido. Los datos de las tablas salen del Terraform de /infra;
si cambia el Terraform, se corrige aquí y se vuelve a generar.
La numeración de tablas y figuras es provisional (continúa la Entrega 1); se ajusta al
ensamblar el informe final, cuando exista la sección 3.
Requiere python-docx.
"""
import copy
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "docs/entregas/Entrega1_Informe_QWallet.docx"
SALIDA = RAIZ / "docs/entregas/Entrega2_Seccion_4_Arquitectura_Controles.docx"
FIGURA = RAIZ / "docs/evidencias/EVD-3-08_arquitectura_flujo_datos.png"

TABLA_INICIAL = 24   # la Entrega 1 termina en la Tabla 23
FIGURA_INICIAL = 7   # y en la Figura 6

doc = docx.Document(str(PLANTILLA))
cuerpo = doc.element.body
sect_final = copy.deepcopy(cuerpo[-1])
for hijo in list(cuerpo):
    cuerpo.remove(hijo)
cuerpo.append(sect_final)
# Página vertical en la última sección (la del documento entero hasta que se agregue otra).
s0 = doc.sections[-1]
s0.orientation = WD_ORIENT.PORTRAIT
s0.page_width, s0.page_height = Emu(7772400), Emu(10058400)

n_tabla = TABLA_INICIAL
n_figura = FIGURA_INICIAL


# ---------------------------------------------------------------- utilidades
def titulo(texto, nivel):
    doc.add_heading(texto, level=nivel)


def parrafo(texto, cursivas=()):
    """Párrafo de cuerpo: justificado, sangría de 1.27 cm. *texto* entre asteriscos va en cursiva."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    partes = texto.split("*")
    for i, parte in enumerate(partes):
        if parte:
            r = p.add_run(parte)
            r.font.size = Pt(12)
            r.italic = i % 2 == 1
    return p


def rotulo(tipo, numero, nombre):
    p = doc.add_paragraph()
    r = p.add_run(f"{tipo} {numero}")
    r.bold = True
    r.font.size = Pt(12)
    p.paragraph_format.keep_with_next = True
    p2 = doc.add_paragraph()
    r2 = p2.add_run(nombre)
    r2.italic = True
    r2.font.size = Pt(12)
    p2.paragraph_format.keep_with_next = True


def nota(texto):
    p = doc.add_paragraph()
    r = p.add_run("Nota. ")
    r.italic = True
    r.font.size = Pt(12)
    r2 = p.add_run(texto)
    r2.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def borde(tc_pr, lados):
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        if lado in lados:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "none")
            e.set(qn("w:sz"), "0")
            e.set(qn("w:color"), "FFFFFF")
        b.append(e)
    tc_pr.append(b)


def margenes(tc_pr):
    m = OxmlElement("w:tcMar")
    for lado, v in (("top", 40), ("left", 90), ("bottom", 40), ("right", 90)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tc_pr.append(m)


FUENTE_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FUENTE_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"


def anchos_sin_cortes(encabezado, filas, tam, total):
    """Anchos (pulgadas) para que ninguna palabra se parta en una celda.

    Mínimo por columna = palabra más ancha (Liberation Sans tiene las mismas medidas que
    Arial) más los márgenes de la celda; lo que sobra se reparte según el largo del texto.
    """
    from PIL import ImageFont
    fr = ImageFont.truetype(FUENTE_R, tam * 10)
    fb = ImageFont.truetype(FUENTE_B, tam * 10)
    n = len(encabezado)
    minimo, peso = [0.0] * n, [0.0] * n
    for i, fila in enumerate([encabezado] + filas):
        f = fb if i == 0 else fr
        for j, v in enumerate(fila):
            for palabra in str(v).replace("\n", " ").split():
                w = f.getlength(palabra) / 10 / 72  # pulgadas
                minimo[j] = max(minimo[j], w)
            peso[j] += fr.getlength(str(v)) / 10 / 72
    margen = 180 / 1440 + 0.03
    minimo = [m + margen for m in minimo]
    libre = total - sum(minimo)
    if libre < 0:
        return None
    return [m + libre * p / sum(peso) for m, p in zip(minimo, peso)]


def tabla(nombre, encabezado, filas, anchos_pulg=None, tam=10, nota_txt=None, total=6.5):
    """Tabla en estilo APA 7: líneas solo arriba y abajo del encabezado y al final."""
    global n_tabla
    anchos_pulg = anchos_sin_cortes(encabezado, filas, tam, total)
    if anchos_pulg is None:
        raise SystemExit(f"La tabla '{nombre}' no cabe sin partir palabras en {total} pulgadas")
    rotulo("Tabla", n_tabla, nombre)
    n_tabla += 1
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezado))
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tbl_pr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos_pulg):
        gc.set(qn("w:w"), str(int(w * 1440)))
    todas = [encabezado] + filas
    for i, fila in enumerate(todas):
        row = t.rows[i]
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, valor in enumerate(fila):
            celda = row.cells[j]
            celda.width = Inches(anchos_pulg[j])
            tc_pr = celda._tc.get_or_add_tcPr()
            lados = set()
            if i == 0:
                lados = {"top", "bottom"}
            elif i == len(todas) - 1:
                lados = {"bottom"}
            borde(tc_pr, lados)
            margenes(tc_pr)
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            lineas = str(valor).split("\n")
            for k, linea in enumerate(lineas):
                if k:
                    p = celda.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # Guion que no corta línea: evita que "AME-06" quede partido en la celda.
                r = p.add_run(linea.replace("-", "\u2011"))
                r.font.size = Pt(tam)
                r.bold = i == 0
    if nota_txt:
        nota(nota_txt)
    return t


def figura(nombre, ruta, ancho_pulg, nota_txt):
    global n_figura
    rotulo("Figura", n_figura, nombre)
    n_figura += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Inches(ancho_pulg))
    nota(nota_txt)


def seccion(horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Emu(10058400), Emu(7772400)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Emu(7772400), Emu(10058400)


# ---------------------------------------------------------------- contenido
titulo("4. Mitigación y ciberseguridad", 1)

titulo("4.1 Rediseño arquitectónico cloud seguro", 2)
parrafo(
    "Como Q-Wallet se diseñó segura desde el inicio, el rediseño no corrige una versión "
    "vulnerable: fija la arquitectura definitiva que planteamos en la sección 1.5 y la "
    "convierte en infraestructura como código. Escribimos toda la arquitectura en Terraform "
    "(carpeta infra del repositorio) y la desplegamos en la región us-east-1 desde GitHub "
    "Actions, sin llaves de acceso guardadas. Así, lo que describe esta sección es lo mismo "
    "que existe en la cuenta de AWS y lo mismo que revisa Checkov en la sección 4.5."
)

titulo("4.1.1 Flujo de datos del cliente a la base de datos", 3)
parrafo(
    "La Figura 7 muestra el recorrido de una solicitud desde la app Android hasta DynamoDB. "
    "Las flechas verdes son datos del usuario y las azules punteadas son verificaciones que "
    "hace la plataforma. Las zonas punteadas son las fronteras de confianza TB1 a TB6 de la "
    "sección 1.5. La Tabla 24 describe cada paso numerado y el control que lo protege."
)
seccion(horizontal=True)
figura(
    "Arquitectura de seguridad cloud de Q-Wallet con el flujo de datos numerado",
    FIGURA, 7.2,
    "Elaboración propia a partir del Terraform del repositorio (carpeta infra). Las flechas "
    "verdes son datos del usuario; las azules punteadas, verificaciones de la plataforma; las "
    "grises, despliegue y registros. El WAF se enciende solo durante pruebas y capturas por costo.",
)
seccion(horizontal=False)
tabla(
    "Pasos del flujo de datos y control aplicado en cada uno",
    ["Paso", "Origen → destino", "Qué viaja", "Cómo se protege", "Control"],
    [
        ["1", "App → Cognito", "Usuario, contraseña y código TOTP",
         "Login administrado de Cognito con OAuth2 código de autorización y PKCE; la app no ve la contraseña", "CTRL-18"],
        ["2", "Cognito → App", "Access token (15 min), ID token y refresh token (12 h, con rotación)",
         "Tokens firmados por Cognito; la app los guarda en el Keystore de Android", "CTRL-07, CTRL-18"],
        ["3", "App → AWS WAF", "Solicitud HTTPS con el access token",
         "Solo TLS 1.3 y pinning de la llave pública; el WAF filtra ataques comunes y limita a 100 solicitudes por IP en 5 minutos", "CTRL-08, CTRL-17"],
        ["4", "WAF → API Gateway", "Solicitud permitida por el WAF",
         "Política SecurityPolicy_TLS13_1_3_2025_09 en modo STRICT; validación de parámetros; throttling de 10 solicitudes por segundo y ráfaga de 20", "CTRL-05, CTRL-08"],
        ["5", "API Gateway ↔ Cognito", "Firma, vigencia y alcance del token",
         "El authorizer rechaza con 401 todo token vencido, alterado o sin el alcance qwallet/api; las 18 rutas lo exigen", "CTRL-18"],
        ["6", "API Gateway → Lambda", "Solicitud y claims del usuario (sub)",
         "Invocación firmada por AWS; cada función acepta solo llamadas de la etapa v1 de esta API", "CTRL-12"],
        ["7", "Lambda → DynamoDB", "Lectura o escritura de los datos del usuario",
         "Rol IAM propio por función con acceso solo a sus tablas; el handler valida que el recurso pertenece al sub del token", "CTRL-01, CTRL-12"],
        ["8", "DynamoDB ↔ KMS", "Llave de datos de la tabla",
         "Cifrado AES-256 con la llave propia alias/qwallet-datos; solo se usa a través de DynamoDB", "CTRL-19"],
        ["9", "Lambda → SSM", "Secretos del servicio",
         "SecureString; cada función lee solo su ruta /qwallet/svc-*/", "CTRL-16"],
    ],
    [0.45, 1.25, 1.45, 2.45, 0.9],
    nota_txt="Elaboración propia. Los pasos 1 a 5 cruzan la frontera TB2 y los pasos 6 a 9, las "
             "fronteras TB3 y TB4 (sección 1.5, Tabla 8).",
)

titulo("4.1.2 Enfoque Zero Trust", 3)
parrafo(
    "Seguimos el principio de confianza cero del NIST: ninguna solicitud se acepta por venir "
    "de una red “interna”; cada acceso se autentica y se autoriza por separado (Rose et al., "
    "2020). En Q-Wallet esto se cumple en cuatro niveles. El usuario prueba su identidad con "
    "contraseña y MFA, y cada solicitud lleva un token de corta vida que API Gateway verifica. "
    "Cada servicio firma sus llamadas a AWS con su propio rol IAM, que solo alcanza sus tablas. "
    "Cada dato en reposo depende de una llave KMS cuyo uso también pasa por IAM. Y el "
    "despliegue entra a AWS con una identidad temporal de GitHub (OIDC) limitada al entorno "
    "dev y a la rama main."
)
parrafo(
    "Por la misma razón, las Lambdas no están dentro de una VPC. Una VPC con salida a internet "
    "exigiría un NAT Gateway, que el presupuesto del proyecto prohíbe, y no agregaría "
    "seguridad real: las funciones solo hablan con servicios administrados por sus endpoints "
    "con TLS y con permisos IAM. La segmentación es lógica, por rol (capa 4 de la sección 2.5)."
)

titulo("4.1.3 Despliegue automatizado con identidad temporal", 3)
parrafo(
    "La única configuración manual es un bootstrap que se corre una vez en AWS CloudShell "
    "(archivo infra/bootstrap/bootstrap.sh): crea el bucket del estado de Terraform, el "
    "proveedor OIDC de GitHub, un límite de permisos y dos roles. Todo lo demás lo crea "
    "Terraform desde GitHub Actions. La Tabla 25 resume los roles y por qué no pueden "
    "escalar privilegios. Además, el flujo se detiene si el plan de Terraform destruye o "
    "reemplaza un recurso que no esté aprobado en un archivo del repositorio; así un error "
    "en el código no borra datos. Esto atiende el plan de RSK-32 de la sección 2.5.5."
)
tabla(
    "Identidades de despliegue y sus límites",
    ["Identidad", "Quién la usa", "Condición para asumirla", "Qué puede hacer", "Qué tiene prohibido"],
    [
        ["qwallet-gha-plan", "GitHub Actions en cada pull request",
         "Token OIDC del repositorio con sub repo:dbarrios95@58736977/q-wallet@1403562245:pull_request",
         "Solo leer la configuración y el estado para calcular el plan", "Escribir cualquier recurso; leer datos de las tablas"],
        ["qwallet-gha-despliegue", "GitHub Actions al hacer merge a main",
         "Token OIDC con sub del entorno dev (el entorno solo acepta la rama main)",
         "Crear y modificar recursos con prefijo qwallet; crear roles solo si llevan el límite de permisos",
         "Borrar la llave KMS, tablas, user pool, buckets, trail o repositorios; detener CloudTrail; tocar los roles de GitHub o el límite"],
        ["qwallet-limite-permisos", "Todo rol que crea Terraform (Lambdas y API Gateway)",
         "Se exige al crear el rol",
         "Techo de permisos: tablas qwallet-*, llave del proyecto, logs, SSM /qwallet/*",
         "Todo lo demás, aunque una política lo permita"],
    ],
    [1.25, 1.15, 1.45, 1.35, 1.3],
    tam=9,
    nota_txt="Elaboración propia a partir de infra/bootstrap/politicas. El límite de permisos "
             "fija el máximo que puede tener un rol; un permiso solo vale si lo dan la política "
             "y el límite a la vez (Amazon Web Services [AWS], s. f.-d).",
)

titulo("4.2 Controles de seguridad propuestos", 2)
parrafo(
    "Esta sección especifica los dos controles que pide la fase de diseño: el menor privilegio "
    "en los roles IAM y la estrategia de cifrado. Cierra con la matriz que une cada amenaza "
    "del modelo STRIDE con su requisito de seguridad y el control que lo cumple."
)

titulo("4.2.1 Menor privilegio en roles IAM", 3)
parrafo(
    "Cada microservicio tiene su propio rol IAM. Para que los permisos se puedan dar tabla "
    "por tabla, separamos los datos en ocho tablas de DynamoDB según su tipo, en lugar de una "
    "sola tabla compartida. La Tabla 26 muestra la matriz de acceso que aplica el Terraform "
    "(archivo infra/iam.tf): ninguna función tiene permiso de Scan, ninguna usa comodines "
    "de acción o de tabla, y svc-budgets no puede leer transferencias ni el PIN."
)
tabla(
    "Matriz de menor privilegio: acciones de cada rol de Lambda por tabla",
    ["Tabla", "svc-profile", "svc-accounts", "svc-transfers", "svc-budgets", "svc-admin"],
    [
        ["perfiles", "Get, Put, Update; Query en índice por-telefono", "—", "Query en índice por-telefono", "—", "—"],
        ["pin", "Put, Update", "—", "Get, Update", "—", "—"],
        ["cuentas", "—", "Get, Query", "Get, Query, Update, ConditionCheck", "Query", "—"],
        ["movimientos", "—", "Query", "Put", "Query", "—"],
        ["transferencias", "—", "—", "Get, Put", "—", "—"],
        ["beneficiarios", "—", "—", "Query, Put, Delete", "—", "—"],
        ["presupuestos", "—", "—", "—", "Get, Query, Put, Update", "—"],
        ["auditoria", "Put", "—", "Put", "—", "Query, Put"],
        ["Otros", "SSM propio, logs propios", "SSM propio, logs propios", "SSM propio, logs propios",
         "SSM propio, logs propios", "Cognito: listar, bloquear y cerrar sesión de usuarios"],
    ],
    [1.05, 1.15, 0.9, 1.25, 1.05, 1.1],
    tam=9,
    nota_txt="Elaboración propia a partir de infra/iam.tf. Get, Put, Update, Delete y Query son "
             "las acciones GetItem, PutItem, UpdateItem, DeleteItem y Query de DynamoDB. Todas las "
             "funciones usan además la llave KMS solo a través de DynamoDB (condición "
             "kms:ViaService) y escriben trazas en X-Ray.",
)
parrafo(
    "El menor privilegio no queda solo escrito. Después de cada despliegue, el flujo "
    "automático consulta el simulador de políticas de IAM con cada rol contra cada tabla y "
    "guarda el resultado como evidencia EVD-3-11. [EVIDENCIA PENDIENTE: se genera en el "
    "primer despliegue.]"
)

titulo("4.2.2 Estrategia de cifrado", 3)
parrafo(
    "En reposo, todos los datos de negocio se cifran con AES-256. DynamoDB usa cifrado de "
    "sobre: la llave raíz de cada tabla queda protegida por una llave de AWS KMS y los datos "
    "se cifran con AES-256 (AWS, s. f.-e). Usamos una llave propia (alias/qwallet-datos) de "
    "tipo SYMMETRIC_DEFAULT, que en AWS KMS corresponde a AES-256-GCM, con rotación "
    "automática cada 365 días (AWS, s. f.-f). La misma llave cifra los logs de CloudWatch, "
    "los registros de CloudTrail y las variables de entorno de las Lambdas."
)
parrafo(
    "En tránsito, el endpoint de API Gateway usa la política SecurityPolicy_TLS13_1_3_2025_09 "
    "en modo STRICT, que acepta solo TLS 1.3 (AWS, s. f.-g). Por eso subimos el mínimo de la "
    "app a Android 10 (minSdk 29), la primera versión que trae TLS 1.3. Este cambio "
    "reemplaza la decisión de la Entrega 1, que dejaba TLS 1.2 de respaldo. La Tabla 27 "
    "resume dónde está cada dato sensible y cómo se protege."
)
tabla(
    "Estrategia de cifrado por tipo de dato",
    ["Dato", "Dónde está", "En reposo", "En tránsito", "Llave y quién la usa"],
    [
        ["Saldos, movimientos y transferencias", "DynamoDB (cuentas, movimientos, transferencias)",
         "AES-256 con la llave propia", "TLS 1.3 app–API; HTTPS firmado Lambda–DynamoDB", "alias/qwallet-datos; solo los roles de sus Lambdas, vía DynamoDB"],
        ["DPI y teléfono", "DynamoDB (perfiles)",
         "AES-256; el teléfono se busca por hash y se muestra enmascarado (CTRL-03)", "Igual que arriba", "Igual que arriba"],
        ["PIN de transacción", "DynamoDB (pin)",
         "Hash scrypt con sal, nunca el PIN; la tabla además va cifrada", "TLS 1.3, solo al autorizar una operación", "No hay llave: el hash no se puede revertir"],
        ["Tokens de sesión", "Teléfono del cliente",
         "Android Keystore con desbloqueo biométrico (CTRL-07)", "TLS 1.3 con pinning", "Keystore del dispositivo"],
        ["Contraseñas y secretos TOTP", "Amazon Cognito", "Los administra Cognito", "TLS del login administrado", "AWS"],
        ["Secretos de los servicios", "SSM Parameter Store", "SecureString cifrado con KMS", "HTTPS firmado", "Cada Lambda solo su ruta"],
        ["Logs de aplicación y de API", "CloudWatch Logs", "AES-256 con la llave propia; 365 días", "Interno de AWS", "Servicio de logs, solo para esta cuenta"],
        ["Registros de auditoría", "S3 (bucket de CloudTrail)", "SSE-KMS con la llave propia; validación de integridad", "HTTPS; el bucket rechaza conexiones sin TLS", "CloudTrail, solo para este trail"],
        ["Estado de Terraform", "S3 (bucket del estado)", "SSE-KMS; versionado; sin acceso público", "HTTPS; rechaza conexiones sin TLS", "Roles de GitHub por OIDC"],
    ],
    [1.25, 1.25, 1.55, 1.3, 1.15],
    tam=9,
    nota_txt="Elaboración propia a partir de infra/kms.tf, dynamodb.tf, cloudtrail.tf y "
             "bootstrap/bootstrap.sh. Las imágenes de contenedor en ECR se cifran con AES-256 "
             "administrado por AWS porque no contienen datos de clientes ni secretos.",
)
parrafo(
    "La prueba de TLS (EVD-3-12) y el reporte de cifrado (EVD-5-04) se generan "
    "automáticamente después de cada despliegue: la conexión con TLS 1.3 debe aceptarse, la "
    "de TLS 1.2 debe rechazarse y cada tabla debe mostrar la llave propia. [EVIDENCIA "
    "PENDIENTE: se genera en el primer despliegue.]"
)

seccion(horizontal=True)
titulo("4.2.3 Matriz de requisitos de seguridad", 3)
parrafo(
    "La Tabla 28 traduce las amenazas del modelo STRIDE (sección 2.3) en requisitos "
    "verificables, tomados de OWASP MASVS v2 para la app y del OWASP API Security Top 10 "
    "para la API, y dice dónde se cumple cada uno. La columna Estado separa lo que ya está "
    "en la infraestructura desplegada de lo que se implementa en el código de los servicios "
    "y de la app en la fase de desarrollo."
)
tabla(
    "Matriz de requisitos de seguridad",
    ["Req.", "Amenaza", "Requisito", "Referencia", "Control y dónde", "Estado"],
    [
        ["REQ-01", "AME-05, AME-06", "Todo el tráfico de la app va cifrado con TLS 1.3 y la app fija la llave del servidor", "MASVS-NETWORK-1, NETWORK-2", "CTRL-08: apigateway.tf y la app", "Infraestructura lista; pinning en la app pendiente"],
        ["REQ-02", "AME-13, AME-16", "Ninguna ruta acepta solicitudes sin un token válido y con el alcance de la API", "API2:2023", "CTRL-18: authorizer en las 18 rutas (apigateway.tf)", "Implementado"],
        ["REQ-03", "AME-04, AME-15", "El inicio de sesión exige MFA y usa código de autorización con PKCE", "MASVS-AUTH-1", "CTRL-18: cognito.tf", "Implementado"],
        ["REQ-04", "AME-38", "Las operaciones de dinero piden un segundo factor (PIN) con bloqueo por intentos", "MASVS-AUTH-3, API6:2023", "CTRL-02: svc-transfers y tabla pin", "Tabla lista; lógica en desarrollo"],
        ["REQ-05", "AME-17, AME-37", "Cada servicio entrega solo recursos del usuario del token", "API1:2023", "CTRL-01: handlers de los servicios", "En desarrollo"],
        ["REQ-06", "AME-19", "Cada función tiene solo los permisos que necesita", "API5:2023; ISO 27001 8.2", "CTRL-12: iam.tf y límite de permisos", "Implementado"],
        ["REQ-07", "AME-26", "Los datos en reposo se cifran con AES-256 y llave propia", "MASVS-STORAGE-1; ISO 27001 8.24", "CTRL-19: kms.tf y dynamodb.tf", "Implementado"],
        ["REQ-08", "AME-27, AME-08", "Ningún secreto vive en el cliente ni en el código", "MASVS-STORAGE-1; CWE-798", "CTRL-16: SSM por servicio", "Implementado (rutas); gitleaks en la fase 3"],
        ["REQ-09", "AME-28, AME-29", "Nadie puede borrar la llave, detener ni borrar la auditoría desde el despliegue", "ISO 27001 8.15", "CTRL-12, CTRL-20: denegaciones del rol de despliegue", "Implementado"],
        ["REQ-10", "AME-30", "Los buckets no son públicos y exigen TLS", "CIS AWS", "CTRL-12: cloudtrail.tf y bootstrap", "Implementado"],
        ["REQ-11", "AME-32", "El despliegue usa identidad temporal y no puede borrar recursos sin aprobación", "SLSA; A08:2025", "CTRL-21: OIDC, entorno dev y revisión del plan", "Implementado"],
        ["REQ-12", "AME-35", "Las imágenes no se pueden reemplazar y su base está fijada", "CIS Docker", "CTRL-11: ecr.tf y Dockerfile", "Implementado"],
        ["REQ-13", "AME-11, AME-18", "La API limita la frecuencia de solicitudes", "API4:2023", "CTRL-05, CTRL-17: throttling y WAF", "Implementado (WAF solo en pruebas)"],
        ["REQ-14", "AME-14, AME-25", "Los errores son genéricos y los logs no guardan tokens ni cuerpos", "API8:2023; A09:2025", "CTRL-09, CTRL-13: respuestas y logs de API Gateway", "Implementado en API Gateway"],
    ],
    tam=9,
    total=9.0,
    nota_txt="Elaboración propia. Las amenazas AME son las del modelo de la sección 2.3 y los "
             "controles CTRL, los de la sección 1.5. MASVS: OWASP Mobile Application Security "
             "Project (2024); API: OWASP API Security Top 10 2023; A08 y A09: OWASP Top 10:2025.",
)
seccion(horizontal=False)

titulo("4.5 Hardening de infraestructura cloud", 2)
parrafo(
    "El endurecimiento de la infraestructura cloud se aplica en el propio Terraform "
    "(secciones 4.1 y 4.2) y se audita de forma automática. "
    "Revisamos el Terraform con Checkov 3.3.22 en cada pull request, antes de que algo llegue "
    "a AWS (evidencia EVD-3-06). El primer análisis completo, del 29 de octubre de 2026, "
    "revisó 125 bloques de Terraform: 282 verificaciones pasaron, 0 fallaron y 39 se omitieron con "
    "justificación. Cada omisión está escrita junto al recurso en el propio código, con el "
    "formato #checkov:skip=ID:motivo, para que el motivo quede a la vista de quien revisa. "
    "El Dockerfile de las Lambdas pasó las 5 verificaciones que le aplican."
)
parrafo(
    "Antes de llegar a ese resultado, la primera corrida marcó 52 fallas. Corregimos las que "
    "tenían arreglo dentro del presupuesto: subimos la retención de logs a 365 días, "
    "activamos el rastreo con X-Ray en API Gateway y en las Lambdas y protegimos la API "
    "contra reemplazos. Las demás se aceptaron por las razones de la Tabla 29."
)
tabla(
    "Verificaciones de Checkov omitidas y su justificación",
    ["Verificación", "Qué pide", "Recursos", "Por qué se omite"],
    [
        ["CKV_AWS_117", "Lambda dentro de una VPC", "5 Lambdas", "Exigiría NAT Gateway (prohibido por costo); segmentación por rol IAM (4.1.2)"],
        ["CKV_AWS_115", "Concurrencia reservada por función", "5 Lambdas", "La cuenta nueva tiene un límite de 10 ejecuciones simultáneas; el freno lo pone el throttling de API Gateway"],
        ["CKV_AWS_116", "Cola de mensajes fallidos", "5 Lambdas", "Invocación síncrona: el error vuelve al cliente"],
        ["CKV_AWS_272", "Firma de código", "5 Lambdas", "No aplica a imágenes de contenedor; la integridad la dan las etiquetas inmutables de ECR"],
        ["CKV_AWS_136", "ECR cifrado con KMS", "5 repositorios", "Las imágenes no tienen datos de clientes ni secretos; AES-256 administrado por AWS"],
        ["CKV_AWS_120, CKV_AWS_225", "Caché de API Gateway", "Etapa v1", "Respuestas con datos de un solo cliente; la caché tiene costo por hora"],
        ["CKV2_AWS_4", "Log de ejecución de API Gateway", "Etapa v1", "Puede copiar cuerpos con datos personales; usamos el log de acceso (CTRL-09)"],
        ["CKV2_AWS_51", "Certificado de cliente", "Etapa v1", "Pensado para backends HTTP; aquí la integración con Lambda la firma AWS"],
        ["CKV2_AWS_29, CKV2_AWS_77", "WAF asociado con regla de Log4j", "Etapa v1", "El WAF (con la regla de entradas maliciosas, que cubre Log4j) se enciende solo en pruebas por costo"],
        ["CKV_AWS_109, CKV_AWS_111, CKV_AWS_356", "Sin “*” en políticas", "Política de la llave KMS", "En una política de llave, “*” significa “esta llave”; los servicios llevan condiciones de origen"],
        ["CKV_AWS_18, CKV_AWS_144, CKV2_AWS_62", "Registro de acceso, replicación y notificaciones del bucket", "Bucket de CloudTrail", "Costo; una sola región aceptada (RSK-01); las alertas se definen con el SIEM"],
        ["CKV2_AWS_10, CKV_AWS_252", "CloudTrail a CloudWatch y a SNS", "Trail", "Se decide con el SIEM en la fase de operación (CTRL-20)"],
    ],
    [1.45, 1.4, 1.0, 2.65],
    tam=9,
    nota_txt="Elaboración propia a partir del resultado de Checkov 3.3.22 sobre la carpeta infra "
             "(Bridgecrew, s. f.). El escaneo de las imágenes con Trivy se agrega en la fase de "
             "desarrollo, cuando existan las imágenes reales de los servicios.",
)

titulo("Referencias de esta sección [se integran a la lista general al ensamblar]", 2)
refs = [
    "Amazon Web Services. (s. f.-d). Permissions boundaries for IAM entities. AWS Identity and Access Management User Guide. Recuperado el 29 de octubre de 2026, de https://docs.aws.amazon.com/IAM/latest/UserGuide/access_policies_boundaries.html",
    "Amazon Web Services. (s. f.-e). DynamoDB encryption at rest: How it works. Amazon DynamoDB Developer Guide. Recuperado el 29 de octubre de 2026, de https://docs.aws.amazon.com/amazondynamodb/latest/developerguide/encryption.howitworks.html",
    "Amazon Web Services. (s. f.-f). Key spec reference. AWS Key Management Service Developer Guide. Recuperado el 29 de octubre de 2026, de https://docs.aws.amazon.com/kms/latest/developerguide/symm-asymm-choose-key-spec.html",
    "Amazon Web Services. (s. f.-g). Supported security policies. Amazon API Gateway Developer Guide. Recuperado el 29 de octubre de 2026, de https://docs.aws.amazon.com/apigateway/latest/developerguide/apigateway-security-policies-list.html",
    "Bridgecrew. (s. f.). Checkov: Policy-as-code for everyone. Recuperado el 29 de octubre de 2026, de https://www.checkov.io/",
    "Rose, S., Borchert, O., Mitchell, S. y Connelly, S. (2020). Zero trust architecture (NIST Special Publication 800-207). National Institute of Standards and Technology. https://doi.org/10.6028/NIST.SP.800-207",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    for r in p.runs:
        r.font.size = Pt(12)

doc.save(str(SALIDA))
print("Generado:", SALIDA)
