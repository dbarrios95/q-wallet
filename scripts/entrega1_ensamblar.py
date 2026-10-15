"""Ensambla el Word final de la Entrega 1 a partir de los Word por sección.

Uso (desde la raíz del repo):
    python scripts/entrega1_ensamblar.py [salida.docx]

Orden: sección 0 -> 1.1-1.4 -> 1.5 -> 2 -> Referencias unificadas.
Aplica las correcciones del 14-oct-2026 (JM-98-2025, OE3/OE4, citas, registro),
quita las listas de referencias de cada sección, normaliza los títulos (APA 7)
y deja el número de página en todas las secciones. El índice se actualiza con
scripts/entrega1_indice.py después de renderizar.
Requiere: python-docx, docxcompose, lxml.
"""
import copy
import os
import re
import sys
import tempfile
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docxcompose.composer import Composer
from lxml import etree

E = 'docs/entregas'
F0 = f'{E}/Entrega1_Seccion_0_Portada_Resumen_Indice.docx'
F14 = f'{E}/Entrega1_Secciones_1-1_a_1-4.docx'
F15 = f'{E}/Entrega1_Seccion_1-5_Scope_Tecnico.docx'
F2 = f'{E}/Entrega1_Seccion_2_Activos_Riesgos.docx'
FR = f'{E}/Entrega1_Referencias.docx'
OUT = sys.argv[1] if len(sys.argv) > 1 else f'{E}/Entrega1_Informe_QWallet.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def rep(x, old, new, n=1):
    c = x.count(old)
    if c != n:
        raise SystemExit(f'Se esperaban {n} coincidencias y hay {c}: {old[:90]!r}')
    return x.replace(old, new)


def edit_docx(src, edits):
    """Aplica reemplazos de texto sobre word/document.xml y devuelve un .docx temporal."""
    zin = zipfile.ZipFile(src)
    x = zin.read('word/document.xml').decode('utf-8')
    for e in edits:
        x = rep(x, *e)
    fd, tmp = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = x.encode('utf-8') if it.filename == 'word/document.xml' else zin.read(it.filename)
            zout.writestr(it, data)
    return tmp


def ptext(p):
    return ''.join(t.text or '' for t in p.iter(qn('w:t')))


def drop_references(doc, starts):
    """Quita el título de referencias de la sección y todo lo que le sigue."""
    body = doc.element.body
    kids = list(body)
    idx = next(i for i, el in enumerate(kids)
               if el.tag == qn('w:p') and ptext(el).strip().startswith(starts))
    for el in kids[idx:]:
        if el.tag != qn('w:sectPr'):
            body.remove(el)


# ---------------------------------------------------------------- 1.1 a 1.4
E14 = [
    # 1.1: norma vigente y cita AWS (s. f.-a)
    ('conforme al Reglamento para la Administración del Riesgo Tecnológico, Resolución JM-104-2021 de la Junta Monetaria (Amazon Web Services [AWS], s.f.; Junta Monetaria, 2021).',
     'conforme al Reglamento para la Administración del Riesgo Tecnológico de la Junta Monetaria, emitido por la Resolución JM-98-2025, que derogó la Resolución JM-104-2021 (Amazon Web Services [AWS], s. f.-a; Carrillo &amp; Asociados, 2025).'),
    ('(AWS, 2026)', '(AWS, s. f.-b)', 3),
    # 1.2: OE3 y OE4
    ('ejecutando los 13 casos de abuso definidos (CA-01 a CA-13) dentro de la política de pruebas de AWS',
     'ejecutando los 13 casos de abuso del catálogo (CA-01 a CA-13) dentro de la política de pruebas de AWS, con el diseño ajustado que resume la Tabla 11,'),
    ('Implementar los 21 controles de diseño definidos (CTRL-01 a CTRL-21)',
     'Implementar los 23 controles de diseño definidos (CTRL-01 a CTRL-23)'),
    # 1.4: Tabla 6 y texto
    ('Resolución JM-104-2021 (Junta Monetaria, 2021)',
     'Resolución JM-98-2025 (Junta Monetaria, 2025, como se citó en Carrillo &amp; Asociados, 2025)'),
    ('La Resolución JM-104-2021 obliga a las entidades bajo la supervisión de la SIB;',
     'La Resolución JM-98-2025, publicada en el Diario de Centro América el 30 de octubre de 2025, emitió un nuevo Reglamento para la Administración del Riesgo Tecnológico y derogó el de la Resolución JM-104-2021 (Carrillo &amp; Asociados, 2025). No pudimos consultar el texto oficial en el sitio del Banco de Guatemala, por lo que la citamos a partir de esa fuente secundaria. Esta resolución obliga a las entidades bajo la supervisión de la SIB;'),
]

# ---------------------------------------------------------------------- 1.5
E15 = [
    ('sobre el que se construye y se evaluará Q-Wallet', 'sobre el que construimos y evaluaremos Q-Wallet'),
    ('Q-Wallet se construye con una arquitectura', 'Construimos Q-Wallet con una arquitectura'),
    ('concentra los controles del lado del cliente:',
     'concentra los controles del lado del cliente, alineados con OWASP MASVS (OWASP Mobile Application Security Project, 2024):'),
    ('La aplicación se construye con todos los controles desde el diseño;',
     'Construimos la aplicación con todos los controles desde el diseño;'),
    ('el rol y los controles que aplican.',
     'el rol y los controles que aplican. La autorización a nivel de objeto (CTRL-01) y por función (CTRL-14) responde a las categorías API1 y API5 del OWASP API Security Top 10 2023 (OWASP Foundation, 2023).'),
    ('El modelo de amenazas identifica seis fronteras',
     'El modelo de amenazas, elaborado en OWASP Threat Dragon (OWASP Foundation, 2024), identifica seis fronteras'),
    ('Esta tabla es la referencia del mapeo que la Figura 1 representa de forma gráfica.',
     'Esta tabla es la referencia del mapeo que la Figura 1 representa de forma gráfica. Los controles del contenedor (CTRL-11) y de la configuración de la cuenta (CTRL-12) siguen los benchmarks de Docker y de AWS del Center for Internet Security (CIS, 2024, 2025), y los de cifrado y registro de eventos se apoyan en el Anexo A de ISO/IEC 27001:2022 (ISO e IEC, 2022a).'),
    ('Las herramientas se integran en el pipeline de GitHub Actions',
     'Integramos las herramientas en el pipeline de GitHub Actions'),
    ('o se ejecutan de forma manual en la fase de evaluación', 'o las ejecutamos de forma manual en la fase de evaluación'),
    ('Cuando una versión aún no está fijada, se marca como [VERIFICAR].',
     'Cuando una herramienta aún no está instalada indicamos «vigente al instalar»: usaremos la versión estable de ese momento y la registraremos en la bitácora del proyecto.'),
    ('>[VERIFICAR]<', '>Vigente al instalar<'),
    ('Las pruebas de intrusión verifican que cada control funciona. Cada caso de abuso se ejecuta y la evidencia',
     'Las pruebas de intrusión verifican que cada control funciona. Seguimos las fases de planificación, descubrimiento, ataque e informe de la guía NIST SP 800-115 (Scarfone et al., 2008) y tomamos los casos de prueba de la OWASP Web Security Testing Guide (OWASP Foundation, s. f.). Ejecutamos cada caso de abuso y la evidencia'),
    ('Las pruebas se ejecutan únicamente sobre', 'Ejecutamos las pruebas únicamente sobre'),
    ('AWS WAF se activa solo en los días de prueba', 'Activamos AWS WAF solo en los días de prueba'),
    ('Las pruebas que dependen del WAF (por ejemplo, el bloqueo de cargas maliciosas) se programan y se documentan dentro de esas ventanas',
     'Las pruebas que dependen del WAF, como el bloqueo de cargas maliciosas de las categorías del OWASP Top 10:2025 (OWASP Foundation, 2025), las programamos y documentamos dentro de esas ventanas'),
    ('La política de pruebas de intrusión de AWS permite', 'La política de pruebas de intrusión de AWS (AWS, s. f.-b) permite'),
    ('se realizan con ráfagas cortas', 'las realizamos con ráfagas cortas'),
    ('se verifica por configuración, no por ataque', 'la verificamos por configuración, no por ataque'),
    ('Cada evidencia se nombra ', 'Nombramos cada evidencia como '),
    ('y, cuando es posible, muestra la fecha y la hora.', 'y, cuando es posible, incluimos en ella la fecha y la hora.'),
    ('Para cada ataque se graba además un video corto', 'Para cada ataque grabamos además un video corto'),
    ('Las evidencias se versionan en el repositorio', 'Versionamos las evidencias en el repositorio'),
    ('bajo el modelo de responsabilidad compartida.', 'bajo el modelo de responsabilidad compartida (AWS, s. f.-c).'),
    ('se incluyen por su relevancia', 'los incluimos por su relevancia'),
]

# ------------------------------------------------------------------------ 2
E2 = [
    ('(International Organization for Standardization [ISO] &amp; International Electrotechnical Commission [IEC], 2022b)',
     '(ISO e IEC, 2022b)'),
    ('ISO &amp; IEC', 'ISO e IEC', 3),
    ('(Amazon Web Services, s. f.)', '(AWS, s. f.-c)'),
    ('MASVS (OWASP Foundation, s. f.)', 'MASVS (OWASP Mobile Application Security Project, 2024)'),
    ('referencia JM-104-2021', 'referencia JM-98-2025', 3),
    ('con referencia a la JM-104-2021 y al Decreto', 'con referencia a la JM-98-2025 y al Decreto'),
]

# ------------------------------------------------------------- Referencias
CARRILLO = ('<w:p><w:pPr><w:ind w:left="720" w:hanging="720"/></w:pPr>'
            '<w:r><w:rPr><w:i w:val="false"/><w:iCs w:val="false"/></w:rPr><w:t xml:space="preserve">Carrillo &amp; Asociados. (2025, 6 de noviembre). </w:t></w:r>'
            '<w:r><w:rPr><w:i/><w:iCs/></w:rPr><w:t xml:space="preserve">Nuevas disposiciones en pro de la seguridad tecnológica</w:t></w:r>'
            '<w:r><w:rPr><w:i w:val="false"/><w:iCs w:val="false"/></w:rPr><w:t xml:space="preserve"> (Boletín n.º 182). https://carrillolaw.com/en/2025/11/06/nuevas-disposiciones-en-pro-de-la-seguridad-tecnologica/</w:t></w:r></w:p>')


def edit_refs(src):
    zin = zipfile.ZipFile(src)
    x = zin.read('word/document.xml').decode('utf-8')
    # quitar JM-104-2021 (ya no se cita) y agregar Carrillo antes de Center for Internet Security
    m = re.search(r'<w:p>(?:(?!<w:p>).)*?Junta Monetaria\. \(2021\)\..*?</w:p>', x, re.S)
    x = x[:m.start()] + x[m.end():]
    i = x.index('Center for Internet Security. (2024)')
    s = x.rfind('<w:p>', 0, i)
    x = x[:s] + CARRILLO + x[s:]
    fd, tmp = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            zout.writestr(it, x.encode('utf-8') if it.filename == 'word/document.xml' else zin.read(it.filename))
    return tmp


# ------------------------------------------------------------ títulos APA 7
HEAD_PPR = {
    'Heading1': '<w:pPr xmlns:w="%s"><w:pStyle w:val="Heading1"/><w:keepNext/><w:keepLines/><w:spacing w:before="240" w:after="120" w:line="360" w:lineRule="auto"/><w:jc w:val="center"/><w:outlineLvl w:val="0"/></w:pPr>' % W,
    'Heading2': '<w:pPr xmlns:w="%s"><w:pStyle w:val="Heading2"/><w:keepNext/><w:keepLines/><w:spacing w:before="240" w:after="120" w:line="360" w:lineRule="auto"/><w:jc w:val="left"/><w:outlineLvl w:val="1"/></w:pPr>' % W,
    'Heading3': '<w:pPr xmlns:w="%s"><w:pStyle w:val="Heading3"/><w:keepNext/><w:keepLines/><w:spacing w:before="240" w:after="120" w:line="360" w:lineRule="auto"/><w:jc w:val="left"/><w:outlineLvl w:val="2"/></w:pPr>' % W,
}
HEAD_STYLE = {
    'Heading1': ('heading 1', 0, False),
    'Heading2': ('heading 2', 1, False),
    'Heading3': ('heading 3', 2, True),
}
NUEVA_PAGINA = ('1. Selección y Planificación', '2. Activos y Riesgos', 'Referencias')


def normalize_headings(doc):
    for p in doc.element.body.iter(qn('w:p')):
        ppr = p.find(qn('w:pPr'))
        st = ppr.find(qn('w:pStyle')) if ppr is not None else None
        if st is None or not st.get(qn('w:val'), '').startswith('Heading'):
            continue
        txt = re.sub(r'\s+', ' ', ptext(p)).strip()
        lvl = st.get(qn('w:val'))
        if re.match(r'^1\.5 ', txt):
            lvl = 'Heading2'      # 1.5 es subsección de 1
        elif re.match(r'^1\.5\.\d', txt):
            lvl = 'Heading3'
        new = etree.fromstring(HEAD_PPR[lvl])
        if txt in NUEVA_PAGINA:
            new.insert(3, etree.Element(qn('w:pageBreakBefore')))
        p.replace(ppr, new)
        runs = [r for r in p if r.tag == qn('w:r')]
        for r in runs[1:]:
            p.remove(r)
        r = runs[0]
        for el in list(r):
            r.remove(el)
        t = etree.SubElement(r, qn('w:t'))
        t.text = txt


def fix_styles(doc):
    styles = doc.styles.element
    # interlineado 1.5 explícito como "auto" (sin lineRule, LibreOffice lo toma como exacto y recorta la Figura 1)
    for sp in styles.iter(qn('w:spacing')):
        if sp.get(qn('w:line')) and not sp.get(qn('w:lineRule')):
            sp.set(qn('w:lineRule'), 'auto')
    for sid, (name, lvl, ital) in HEAD_STYLE.items():
        for s in styles.findall(qn('w:style')):
            if s.get(qn('w:styleId')) == sid:
                styles.remove(s)
        xml = ('<w:style xmlns:w="%s" w:type="paragraph" w:styleId="%s"><w:name w:val="%s"/>'
               '<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/>'
               '<w:pPr><w:keepNext/><w:keepLines/><w:spacing w:before="240" w:after="120" w:line="360" w:lineRule="auto"/>'
               '<w:jc w:val="%s"/><w:outlineLvl w:val="%d"/></w:pPr>'
               '<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/><w:b/><w:bCs/>%s'
               '<w:color w:val="000000"/><w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr></w:style>'
               % (W, sid, name, 'center' if lvl == 0 else 'left', lvl,
                  '<w:i/><w:iCs/>' if ital else '<w:i w:val="0"/><w:iCs w:val="0"/>'))
        styles.append(etree.fromstring(xml))


def header_everywhere(doc):
    body = doc.element.body
    final = body.find(qn('w:sectPr'))
    href = final.find(qn('w:headerReference'))
    for sp in body.iter(qn('w:sectPr')):
        if sp is final:
            continue
        for old in sp.findall(qn('w:headerReference')) + sp.findall(qn('w:footerReference')):
            sp.remove(old)
        sp.insert(0, copy.deepcopy(href))


ORDEN = {
    'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
              'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
              'tblCaption', 'tblDescription'],
    'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
             'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
    'sectPr': ['headerReference', 'footerReference', 'footnotePr', 'endnotePr', 'type', 'pgSz', 'pgMar',
               'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign', 'noEndnote',
               'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings'],
}


def schema_order(doc):
    """Ordena los hijos de tblPr, tcPr y sectPr según el esquema (errores heredados del Word de la sección 2)."""
    for tag, order in ORDEN.items():
        key = {k: i for i, k in enumerate(order)}
        for el in doc.element.body.iter(qn('w:' + tag)):
            kids = sorted(list(el), key=lambda k: key.get(etree.QName(k).localname, 999))
            vistos = set()
            for k in kids:
                n = etree.QName(k).localname
                if n in vistos and n not in ('headerReference', 'footerReference'):
                    el.remove(k)
                    continue
                vistos.add(n)
                el.remove(k)
                el.append(k)


def main():
    master = Document(F0)
    comp = Composer(master)
    for src, edits, starts in [(F14, E14, 'Referencias (secciones'),
                               (F15, E15, 'Referencias de la sección 1.5'),
                               (F2, E2, 'Referencias de la sección 2')]:
        d = Document(edit_docx(src, edits))
        drop_references(d, starts)
        comp.append(d)
    comp.append(Document(edit_refs(FR)))
    normalize_headings(master)
    fix_styles(master)
    header_everywhere(master)
    schema_order(master)
    for sp in master.element.body.iter(qn('w:spacing')):
        if sp.get(qn('w:line')) and not sp.get(qn('w:lineRule')):
            sp.set(qn('w:lineRule'), 'auto')
    master.core_properties.title = 'Q-Wallet: Entrega 1 — Selección y Planificación; Activos y Riesgos'
    master.save(OUT)
    print('ok', OUT)


if __name__ == '__main__':
    main()
