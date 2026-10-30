"""Genera las secciones 3.4, 3.5, 4.3, 4.4 y 4.6 del informe (Fase 3, bloque 3).

Uso:
    python scripts/bloque3_diagrama_devsecops.py      # figura EVD-3-16
    python scripts/bloque3_word.py                    # Word de las secciones
    python scripts/entrega1_estilos.py docs/entregas/Final_Secciones_3-4_3-5_4-3_4-4_4-6_DevSecOps.docx

Mismo formato que la Entrega 2 (estilos del informe de la Entrega 1). Las cifras salen de
las ejecuciones de GitHub Actions del PR #36 (comentarios del PR) y de los archivos del
repositorio; si cambian, se corrigen aquí y se vuelve a generar.
Numeración provisional: continúa la sección 4 de la Entrega 2 (Tabla 29, Figura 7) y se
ajusta al ensamblar el informe final. Requiere python-docx, Pillow.
"""
import copy
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "docs/entregas/Entrega1_Informe_QWallet.docx"
SALIDA = RAIZ / "docs/entregas/Final_Secciones_3-4_3-5_4-3_4-4_4-6_DevSecOps.docx"
FIGURA = RAIZ / "docs/evidencias/EVD-3-16_flujo_devsecops.png"

TABLA_INICIAL = 30   # la sección 4 de la Entrega 2 termina en la Tabla 29
FIGURA_INICIAL = 8   # y en la Figura 7

doc = docx.Document(str(PLANTILLA))
cuerpo = doc.element.body
sect_final = copy.deepcopy(cuerpo[-1])
for hijo in list(cuerpo):
    cuerpo.remove(hijo)
cuerpo.append(sect_final)
# Página vertical en la última sección (la del documento entero hasta que se agregue otra).
s0 = doc.sections[-1]
s0.orientation = WD_ORIENT.PORTRAIT
s0.page_width, s0.page_height = Emu(7772400), Emu(10058400)

n_tabla = TABLA_INICIAL
n_figura = FIGURA_INICIAL


# ---------------------------------------------------------------- utilidades
def titulo(texto, nivel):
    doc.add_heading(texto, level=nivel)


def parrafo(texto, cursivas=()):
    """Párrafo de cuerpo: justificado, sangría de 1.27 cm. *texto* entre asteriscos va en cursiva."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    partes = texto.split("*")
    for i, parte in enumerate(partes):
        if parte:
            r = p.add_run(parte)
            r.font.size = Pt(12)
            r.italic = i % 2 == 1
    return p


def rotulo(tipo, numero, nombre):
    p = doc.add_paragraph()
    r = p.add_run(f"{tipo} {numero}")
    r.bold = True
    r.font.size = Pt(12)
    p.paragraph_format.keep_with_next = True
    p2 = doc.add_paragraph()
    r2 = p2.add_run(nombre)
    r2.italic = True
    r2.font.size = Pt(12)
    p2.paragraph_format.keep_with_next = True


def nota(texto):
    p = doc.add_paragraph()
    r = p.add_run("Nota. ")
    r.italic = True
    r.font.size = Pt(12)
    r2 = p.add_run(texto)
    r2.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def borde(tc_pr, lados):
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        if lado in lados:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "none")
            e.set(qn("w:sz"), "0")
            e.set(qn("w:color"), "FFFFFF")
        b.append(e)
    tc_pr.append(b)


def margenes(tc_pr):
    m = OxmlElement("w:tcMar")
    for lado, v in (("top", 40), ("left", 90), ("bottom", 40), ("right", 90)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tc_pr.append(m)


FUENTE_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FUENTE_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"


def anchos_sin_cortes(encabezado, filas, tam, total):
    """Anchos (pulgadas) para que ninguna palabra se parta en una celda.

    Mínimo por columna = palabra más ancha (Liberation Sans tiene las mismas medidas que
    Arial) más los márgenes de la celda; lo que sobra se reparte según el largo del texto.
    """
    from PIL import ImageFont
    fr = ImageFont.truetype(FUENTE_R, tam * 10)
    fb = ImageFont.truetype(FUENTE_B, tam * 10)
    n = len(encabezado)
    minimo, peso = [0.0] * n, [0.0] * n
    for i, fila in enumerate([encabezado] + filas):
        f = fb if i == 0 else fr
        for j, v in enumerate(fila):
            for palabra in str(v).replace("\n", " ").split():
                w = f.getlength(palabra) / 10 / 72  # pulgadas
                minimo[j] = max(minimo[j], w)
            peso[j] += fr.getlength(str(v)) / 10 / 72
    margen = 180 / 1440 + 0.03
    minimo = [m + margen for m in minimo]
    libre = total - sum(minimo)
    if libre < 0:
        return None
    return [m + libre * p / sum(peso) for m, p in zip(minimo, peso)]


def tabla(nombre, encabezado, filas, anchos_pulg=None, tam=10, nota_txt=None, total=6.5):
    """Tabla en estilo APA 7: líneas solo arriba y abajo del encabezado y al final."""
    global n_tabla
    anchos_pulg = anchos_sin_cortes(encabezado, filas, tam, total)
    if anchos_pulg is None:
        raise SystemExit(f"La tabla '{nombre}' no cabe sin partir palabras en {total} pulgadas")
    rotulo("Tabla", n_tabla, nombre)
    n_tabla += 1
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezado))
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    # t.autofit = False ya deja w:tblLayout fijo; no se agrega otro (el esquema lo rechaza).
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos_pulg):
        gc.set(qn("w:w"), str(int(w * 1440)))
    todas = [encabezado] + filas
    for i, fila in enumerate(todas):
        row = t.rows[i]
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, valor in enumerate(fila):
            celda = row.cells[j]
            celda.width = Inches(anchos_pulg[j])
            tc_pr = celda._tc.get_or_add_tcPr()
            lados = set()
            if i == 0:
                lados = {"top", "bottom"}
            elif i == len(todas) - 1:
                lados = {"bottom"}
            borde(tc_pr, lados)
            margenes(tc_pr)
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            lineas = str(valor).split("\n")
            for k, linea in enumerate(lineas):
                if k:
                    p = celda.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # Guion que no corta línea: evita que "AME-06" quede partido en la celda.
                r = p.add_run(linea.replace("-", "\u2011"))
                r.font.size = Pt(tam)
                r.bold = i == 0
    if nota_txt:
        nota(nota_txt)
    return t


def figura(nombre, ruta, ancho_pulg, nota_txt):
    global n_figura
    rotulo("Figura", n_figura, nombre)
    n_figura += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Inches(ancho_pulg))
    nota(nota_txt)


def seccion(horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Emu(10058400), Emu(7772400)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Emu(7772400), Emu(10058400)


# ---------------------------------------------------------------- contenido
titulo("3. Evaluación de seguridad", 1)
parrafo(
    "Esta parte del informe reúne las pruebas que hicimos sobre el código y sus dependencias "
    "durante la Fase 3 del enunciado (desarrollo seguro y automatización DevSecOps). Las "
    "secciones 3.4 y 3.5 presentan el análisis estático de código (SAST) y el análisis de "
    "composición de software (SCA). Las pruebas de penetración de caja negra y gris, la "
    "evaluación OWASP API Top 10 y la evaluación de la arquitectura cloud (secciones 3.1, 3.2, "
    "3.3 y 3.6) corresponden a la Fase 4 y se agregan al informe cuando terminen."
)

titulo("3.4 SAST (código móvil y backend)", 2)
parrafo(
    "Analizamos el código fuente de la app Android (carpeta mobile) y de los cinco "
    "microservicios (carpeta services) con dos herramientas distintas, para que una cubra lo "
    "que la otra no ve. Semgrep busca patrones con reglas públicas de la comunidad y con "
    "reglas propias que escribimos para Q-Wallet. CodeQL, el motor de GitHub, compila el "
    "código a una base de datos y sigue el recorrido de los datos desde la entrada hasta "
    "funciones peligrosas (análisis de flujo). Ambas corren en GitHub Actions en cada pull "
    "request, en cada cambio de la rama main y cada lunes; el resultado se publica en la "
    "pestaña de seguridad del repositorio y como comentario del pull request. La Tabla 30 "
    "resume la última ejecución sobre el código del bloque 3."
)
tabla(
    "Resultado del análisis estático de código (SAST)",
    ["Herramienta", "Alcance", "Reglas", "Hallazgos", "Graves", "Compuerta"],
    [
        ["Semgrep 1.179.0", "services y mobile (TypeScript, TSX, Dockerfile)", "608 públicas + 7 propias", "0", "0", "Pasa"],
        ["CodeQL (security-extended)", "services y mobile (JavaScript y TypeScript)", "103", "0", "0", "Pasa"],
    ],
    nota_txt=(
        "Grave = nivel error o severidad de seguridad de 7.0 o más (escala CVSS que usan "
        "ambas herramientas); un hallazgo grave detiene el flujo. Reglas públicas de Semgrep: "
        "p/owasp-top-ten, p/nodejs, p/typescript, p/javascript, p/react, p/secrets y "
        "p/dockerfile. Ejecución del 30 de octubre de 2026 sobre el pull request 36 "
        "(evidencias EVD-3-02 y EVD-3-03)."
    ),
)
parrafo(
    "Cero hallazgos no significa que el código no tenga errores: significa que ninguna regla "
    "encontró un patrón inseguro conocido. Por eso no nos quedamos con ese resultado. Primero, "
    "escribimos siete reglas propias que convierten controles de diseño en verificaciones "
    "automáticas (Tabla 31): si alguien vuelve a escribir, por ejemplo, un Scan de DynamoDB o "
    "un console.log del cuerpo de la solicitud, el flujo falla aunque las reglas públicas no lo "
    "marquen. Segundo, para comprobar que esas reglas sí detectan lo que dicen, el flujo corre "
    "semgrep --test sobre un archivo con casos buenos y malos; las siete pasaron. Tercero, el "
    "SAST no entiende la lógica de negocio (por ejemplo, si un usuario puede leer la cuenta de "
    "otro). Eso lo cubren las pruebas unitarias por control de la sección 4.4 y, desde afuera, "
    "las pruebas de caja gris de la Fase 4."
)
tabla(
    "Reglas propias de Semgrep para Q-Wallet",
    ["Regla", "Qué detecta", "Control"],
    [
        ["qwallet-dynamodb-scan", "Uso de ScanCommand (lee toda una tabla)", "CTRL-05"],
        ["qwallet-log-cuerpo-solicitud", "console.* con el cuerpo o las cabeceras de la solicitud", "CTRL-09"],
        ["qwallet-comparacion-pin-directa", "Comparar el PIN o su hash con ===", "CTRL-02"],
        ["qwallet-aleatorio-inseguro", "Math.random() en el backend", "CTRL-06"],
        ["qwallet-cuerpo-sin-zod", "JSON.parse del cuerpo sin esquema de validación", "CTRL-04"],
        ["qwallet-zod-sin-strict", "Esquema z.object sin .strict()", "CTRL-04"],
        ["qwallet-access-control-allow-origin", "Cabecera CORS en las respuestas de la API", "CTRL-13"],
    ],
    nota_txt=(
        "Archivo .semgrep/qwallet.yml del repositorio; sus casos de prueba están en "
        ".semgrep/qwallet.ts. Todas son de nivel error salvo qwallet-zod-sin-strict (aviso)."
    ),
)

titulo("3.5 SCA (dependencias móviles y backend)", 2)
parrafo(
    "El análisis de composición revisa las bibliotecas de terceros contra bases de "
    "vulnerabilidades públicas. Revisamos tres cosas: las dependencias de los microservicios, "
    "las de la app Android y las imágenes de contenedor de las Lambdas. Para las dos primeras "
    "usamos npm audit (base de avisos de GitHub) con una compuerta propia y Trivy como segunda "
    "opinión; para las imágenes usamos Trivy. La compuerta (CTRL-10) falla con cualquier aviso "
    "High o Critical que no tenga una excepción aprobada y vigente. Cada excepción vive en el "
    "repositorio con su motivo y una fecha límite; al vencer, el flujo vuelve a fallar. La "
    "Tabla 32 muestra los resultados."
)
tabla(
    "Resultado del análisis de composición de software (SCA)",
    ["Componente", "Herramienta", "Crítica", "Alta", "Media", "Baja", "Compuerta"],
    [
        ["Microservicios (services)", "npm audit / Trivy 0.75.0", "0", "0", "0", "0", "Pasa"],
        ["App Android (mobile), antes", "npm audit", "0", "7", "4", "5", "Falla (zod 4.6.5)"],
        ["App Android (mobile), después", "npm audit", "0", "7", "4", "5", "Pasa con 3 excepciones"],
        ["Imágenes de Lambda, antes", "Trivy 0.75.0", "0", "24", "28", "4", "Falla"],
        ["Imágenes de Lambda, después", "Trivy 0.75.0", "0", "0", "1", "0", "Pasa"],
    ],
    nota_txt=(
        "npm audit cuenta paquetes afectados; en la app, los 7 paquetes de severidad alta "
        "vienen de 3 avisos distintos (Tabla 33). En las imágenes las cifras son por imagen "
        "(las cinco dan lo mismo) y la compuerta solo bloquea las altas o críticas que ya "
        "tienen corrección. La única media que queda es ca-certificates del sistema operativo "
        "base (Amazon Linux 2023), que se corrige cuando AWS publique la imagen nueva. Trivy, "
        "como segunda opinión sobre las dependencias, coincide: 0 en los microservicios y 5 en "
        "la app (2 altas: decode-uri-component y node-forge, las mismas de HAL-02 y HAL-03). "
        "Evidencias EVD-3-04 y EVD-3-05."
    ),
)
parrafo(
    "Todo hallazgo real se registró con un identificador HAL y su tratamiento (Tabla 33). "
    "El más relevante fue HAL-01: la versión de zod que fijamos el 3 de octubre (4.6.5) "
    "está afectada por el aviso GHSA-mrhx-c7xc-m3w5 (agotamiento de memoria al validar listas "
    "muy grandes), publicado el 1 de octubre; no lo notamos al fijarla y apareció al correr "
    "npm audit para armar la compuerta. Es el caso que el análisis automático evita. La subimos a "
    "4.6.7 en la app y en los microservicios, y además todos nuestros esquemas limitan el "
    "tamaño del cuerpo y de las listas. Para demostrar que la compuerta funciona, abrimos un "
    "pull request de prueba (número 37) que volvía a poner zod 4.6.5 en los microservicios: "
    "la etapa SCA falló con ese aviso y el pull request quedó bloqueado (evidencia EVD-3-10)."
)
tabla(
    "Hallazgos de SAST, SCA e imágenes y su tratamiento",
    ["HAL", "Hallazgo", "Severidad", "Tratamiento", "Estado"],
    [
        ["HAL-01", "zod 4.6.5 en la app y en la versión fijada para los microservicios (GHSA-mrhx-c7xc-m3w5)", "Alta", "Mitigar: zod 4.6.7", "Cerrado"],
        ["HAL-02", "decode-uri-component 0.2.2 dentro de expo-router 57 (GHSA-qr28-fpjm-wq7f): congela la app con un enlace malformado", "Alta", "Aceptar hasta el 28-nov-2026: corregirlo exige Expo SDK 58", "Aceptado"],
        ["HAL-03", "node-forge 1.4.0 y zod 3.25.76 dentro de la herramienta de compilación de Expo (@expo/cli)", "Alta", "Aceptar: no van dentro del APK", "Aceptado"],
        ["HAL-04", "La pantalla de bienvenida decía «cifrado de extremo a extremo» y «cumplimiento JM-104-2021»", "Baja", "Mitigar: texto corregido", "Cerrado"],
        ["HAL-05", "Imagen base de Lambda con 24 vulnerabilidades altas corregibles en npm, el emulador local y el SDK que trae el runtime", "Alta", "Mitigar: digest vigente y se quitan esos componentes", "Cerrado"],
    ],
    tam=9,
    nota_txt=(
        "Las excepciones de HAL-02 y HAL-03 están en mobile/excepciones-sca.json, valen solo "
        "para la copia indirecta del paquete, las aprobó el dueño del riesgo y vencen el 28 de "
        "noviembre de 2026. HAL-04 se "
        "encontró por revisión manual del código. El registro completo está en "
        "docs/REGISTRO_RIESGOS.md, sección D.2."
    ),
)
parrafo(
    "OWASP Dependency-Check y Snyk, que sugiere el enunciado, quedaron integrados en el flujo "
    "pero solo corren si existen sus llaves de acceso (NVD_API_KEY y SNYK_TOKEN), que todavía "
    "no creamos. La cobertura actual no depende de ellos: npm audit y Trivy consultan la base "
    "de avisos de GitHub y la de Aqua Security, y Dependabot vigila el repositorio todos los "
    "días."
)

titulo("4.3 Hardening de aplicación móvil", 2)
parrafo(
    "La app Android todavía muestra datos simulados: en este bloque reforzamos su "
    "configuración de compilación y dejamos claro qué controles del cliente faltan. La "
    "Tabla 34 separa lo que ya está en el código, lo que configuramos ahora y lo pendiente. El "
    "cambio principal es subir la versión mínima de Android a la 10 (API 29), que trae TLS 1.3 "
    "de fábrica: así la app puede hablar con la API, que solo acepta TLS 1.3 (sección 4.2.2)."
)
tabla(
    "Controles de hardening de la app Android",
    ["Control", "MASVS", "Cómo", "Estado"],
    [
        ["Android 10 o superior (minSdk 29)", "NETWORK", "expo-build-properties en app.json", "Configurado (bloque 3)"],
        ["Sin tráfico HTTP en claro", "NETWORK", "usesCleartextTraffic: false", "Configurado (bloque 3)"],
        ["Ofuscación y reducción con R8 en la versión de release", "RESILIENCE", "enableMinifyInReleaseBuilds y enableShrinkResourcesInReleaseBuilds", "Configurado (bloque 3)"],
        ["Inspector de red de desarrollo apagado", "CODE", "networkInspector: false", "Configurado (bloque 3)"],
        ["Sin respaldo de datos de la app", "STORAGE", "android.allowBackup: false", "Implementado"],
        ["Bloqueo de capturas de pantalla (CTRL-22)", "PLATFORM", "FLAG_SECURE en transferir, desbloquear y perfil", "Implementado"],
        ["Validación de formularios con zod estricto", "CODE", "login y transferencia; zod 4.6.7 (HAL-01)", "Implementado"],
        ["Textos de seguridad veraces", "PRIVACY", "pantalla de bienvenida (HAL-04)", "Corregido"],
        ["OAuth2 con PKCE y MFA (CTRL-18)", "AUTH", "expo-auth-session contra Cognito", "Pendiente: al conectar la app"],
        ["Tokens en el Keystore con biometría (CTRL-07)", "STORAGE", "expo-secure-store con requireAuthentication", "Pendiente: al conectar la app"],
        ["Pinning de llaves públicas (CTRL-08)", "NETWORK", "react-native-ssl-public-key-pinning", "Pendiente: al conectar la app"],
        ["Detección de root, hooking y emulador (CTRL-15)", "RESILIENCE", "freeRASP (plugin ya registrado)", "Pendiente: sección 5.2"],
    ],
    tam=9,
    nota_txt=(
        "MASVS = OWASP Mobile Application Security Verification Standard 2.1.0. Lo "
        "configurado en el bloque 3 se aplica al generar el proyecto nativo; falta comprobarlo "
        "con una compilación de release en el equipo de desarrollo y con MobSF (evidencia "
        "EVD-3-07)."
    ),
)

titulo("4.4 Hardening de APIs y microservicios", 2)
parrafo(
    "Reemplazamos la imagen provisional de la Entrega 2 por el código real de los cinco "
    "microservicios, escrito en TypeScript estricto. Todos comparten un módulo común que "
    "aplica los mismos controles a cada solicitud, en este orden: verificar el token, validar "
    "la entrada, comprobar que el recurso pertenece al usuario, registrar sin datos sensibles y "
    "responder sin detalles internos. Cada control tiene pruebas automáticas que corren en el "
    "flujo DevSecOps: 60 pruebas en total, todas en verde (evidencia EVD-3-14)."
)
titulo("4.4.1 Prácticas de codificación segura", 3)
parrafo(
    "El enunciado pide demostrar tres prácticas: consultas preparadas contra inyección SQL, "
    "sanitización contra XSS y manejo seguro de tokens JWT. La Tabla 35 muestra cómo se "
    "aplican en Q-Wallet, junto con otras prácticas del mismo tipo."
)
tabla(
    "Prácticas de codificación segura en los microservicios",
    ["Práctica", "Cómo la aplicamos", "Prueba"],
    [
        ["Consultas preparadas (inyección)", "DynamoDB no usa SQL; todas las consultas pasan los valores aparte con ExpressionAttributeValues y la llave siempre lleva el sub del token, nunca un valor armado con texto del cliente", "Día de auditoría con «' OR 1=1» responde 400"],
        ["Sanitización contra XSS", "Esquemas zod con lista de caracteres permitidos en nombres y descripciones; respuestas solo en JSON con Content-Type, X-Content-Type-Options y Content-Security-Policy", "Nombre con <script> responde 400"],
        ["Manejo seguro del JWT", "Además del authorizer, la Lambda verifica el access token con aws-jwt-verify: firma, emisor, vencimiento, token_use, client_id y scope; el sub debe coincidir con el del authorizer", "Firma alterada o sub distinto responde 401"],
        ["Mass assignment", "Esquemas .strict(): un campo no previsto (tienePin, userId, saldo) se rechaza", "PATCH /me con saldoCentavos responde 400"],
        ["Montos exactos", "Centavos enteros positivos con tope; nunca decimales de punto flotante", "Monto -100, 10.5 o «100» responde 400"],
        ["Errores sin detalles", "Mensaje fijo por código y correlationId; la pila y el mensaje interno solo van al log", "Error interno no expone la tabla ni la llave"],
        ["Logs sin datos sensibles", "Redacción automática de PIN, tokens, teléfono, DPI y correo, incluso anidados", "El PIN y el JWT no aparecen en el log"],
        ["Aleatoriedad segura", "randomUUID y randomBytes de Node crypto", "Regla qwallet-aleatorio-inseguro"],
    ],
    tam=9,
    nota_txt="Código en services/comun/src; pruebas en services/*/test (vitest 5.0.3).",
)
titulo("4.4.2 Controles por microservicio", 3)
parrafo(
    "La Tabla 36 resume los controles de cada microservicio y las pruebas que los comprueban. "
    "Las pruebas usan repositorios en memoria con la misma lógica de condiciones que DynamoDB, "
    "así se pueden simular ataques (por ejemplo, veinte intentos de PIN al mismo tiempo) sin "
    "tocar AWS. Las mismas comprobaciones se repetirán contra la API desplegada en la Fase 4."
)
tabla(
    "Controles y pruebas de cada microservicio",
    ["Servicio", "Controles", "Qué comprueban las pruebas", "Pruebas"],
    [
        ["Común (todas)", "CTRL-04, 05, 09, 13, 18", "Token sin firma válida, mass assignment, cuerpo de más de 8 KB, cursor apuntando a otra cuenta, redacción de logs, cabeceras", "17"],
        ["svc-profile", "CTRL-01, 02, 03, 04", "Perfil creado con el sub del token, teléfono enmascarado, teléfono ajeno rechazado, PIN guardado solo como hash scrypt, PIN trivial rechazado", "9"],
        ["svc-accounts", "CTRL-01, 03, 05", "Cuenta ajena responde 404 (CA-01), límite por defecto 20 y máximo 50, cursor manipulado rechazado (CA-05)", "6"],
        ["svc-transfers", "CTRL-01, 02, 04, 06", "Idempotencia, diez solicitudes simultáneas con la misma clave aplican una sola transferencia (CA-06), saldo insuficiente, bloqueo de 15 min al quinto PIN fallido aun con veinte intentos simultáneos (CA-02), transferencia ajena 404", "17"],
        ["svc-budgets", "CTRL-01, 04", "Presupuestos de otro usuario invisibles e inmodificables, categorías cerradas, alertas al 80 % y al 100 %", "5"],
        ["svc-admin", "CTRL-14, 23, 03, 05", "Cliente recibe 403 en /admin y queda auditado (CA-09), grupo parecido no sirve, bloqueo exige motivo y canal de verificación, no se bloquea a otro admin", "6"],
    ],
    tam=9,
    nota_txt=(
        "CA = caso de abuso del registro de riesgos. Ejecución del 30 de octubre de 2026: 60 "
        "de 60 pruebas pasan (EVD-3-14). Tope por transferencia: Q5,000.00 de dinero simulado."
    ),
)
titulo("4.4.3 Imagen de contenedor de cada microservicio", 3)
parrafo(
    "Cada microservicio tiene su Dockerfile de dos etapas (CTRL-11). La primera instala las "
    "dependencias exactas del package-lock y empaqueta el servicio en un solo archivo; la "
    "segunda parte de la imagen oficial de Lambda para Node.js 24 fijada por su digest, copia "
    "solo ese archivo como de solo lectura y corre con un usuario sin privilegios (1000). Trivy "
    "encontró en la imagen base 24 vulnerabilidades altas con corrección (HAL-05), todas en "
    "componentes que Lambda no usa para correr nuestro código: npm, corepack, el emulador local "
    "de Lambda y la copia del SDK de AWS que trae el runtime (nuestro código lleva su propio "
    "SDK empaquetado). Actualizamos el digest y quitamos esos componentes de la imagen final; "
    "Trivy bajó a cero vulnerabilidades altas. Como quitar archivos de la imagen base puede "
    "romper el arranque, agregamos una prueba de humo: el flujo arranca el runtime real de cada "
    "imagen frente a un simulador local de la Runtime API de Lambda, le envía GET /me sin token "
    "y exige la respuesta 401 con las cabeceras de seguridad (EVD-3-15)."
)

titulo("4.6 Implementación de controles DevSecOps", 2)
parrafo(
    "La Figura 8 muestra el flujo completo desde que se abre un pull request hasta que el "
    "cambio queda desplegado y verificado en AWS. La rama main está protegida: nada llega a "
    "producción sin pasar por un pull request, sus compuertas automáticas y la aprobación por "
    "merge. Son dos flujos de GitHub Actions: devsecops.yml revisa el código (Fase 3) e "
    "infra.yml revisa la infraestructura y los contenedores, despliega y verifica (Fase 5)."
)
seccion(horizontal=True)
figura(
    "Flujo DevSecOps de Q-Wallet: compuertas del pull request, despliegue y verificación",
    FIGURA,
    9.0,
    "Elaboración propia (EVD-3-16). Las cajas rojas y verdes marcan dónde se detiene el flujo "
    "y dónde termina. Herramientas descargadas por versión exacta y verificadas por suma "
    "SHA-256; acciones de GitHub fijadas por el hash del commit.",
)
seccion(horizontal=False)
parrafo(
    "La Tabla 37 detalla cada etapa. Tres decisiones hacen que el flujo sea un control y no "
    "solo un reporte. Primera, cada etapa tiene una compuerta que detiene el pull request. "
    "Segunda, las excepciones son explícitas, con motivo y fecha de vencimiento, y no pueden "
    "tapar una dependencia directa. Tercera, el mismo flujo corre cada lunes sobre main, así "
    "una vulnerabilidad publicada después del despliegue aparece sin esperar un cambio de código "
    "(gestión continua, sección 5.5)."
)
tabla(
    "Etapas del flujo DevSecOps y sus compuertas",
    ["Etapa", "Herramienta y versión", "Compuerta", "Evidencia"],
    [
        ["Pruebas por control", "vitest 5.0.3, TypeScript 6.0.3", "Una prueba roja o un error de tipos", "EVD-3-14"],
        ["SAST", "Semgrep 1.179.0", "Hallazgo grave o regla propia que no pasa su prueba", "EVD-3-02"],
        ["SAST", "CodeQL (security-extended)", "Hallazgo grave", "EVD-3-03"],
        ["SCA", "npm audit + scripts/sca_compuerta.py", "Aviso alto o crítico sin excepción vigente", "EVD-3-04, EVD-3-10"],
        ["SCA e inventario", "Trivy 0.75.0 (SBOM CycloneDX)", "Informativa", "EVD-3-04"],
        ["Secretos", "gitleaks 8.30.1 (todo el historial: 89 commits, sin hallazgos)", "Cualquier secreto", "EVD-3-09"],
        ["IaC", "Checkov 3.3.22 (Terraform: 282 pasan; Dockerfile: 295 pasan; 0 fallan)", "Control fallido sin excepción en el código", "EVD-3-06"],
        ["Contenedores", "Docker multi-etapa + prueba de humo", "Imagen que no arranca o no responde 401", "EVD-3-15"],
        ["Contenedores", "Trivy 0.75.0", "Vulnerabilidad alta o crítica con corrección", "EVD-3-05"],
        ["Despliegue", "Terraform 1.16.5 con OIDC", "Plan que borra o reemplaza algo no aprobado", "EVD-3-11"],
        ["Verificación", "verificar_aws.sh", "TLS, menor privilegio o cifrado distintos de lo diseñado", "EVD-3-12, EVD-3-13, EVD-5-04"],
        ["SCA opcional", "Dependency-Check 12.1.0, Snyk 1.1307.4", "CVSS 7 o más / severidad alta", "Requieren llaves"],
    ],
    tam=9,
    nota_txt=(
        "Las evidencias se publican como comentario del pull request y, para main y la revisión "
        "semanal, en los issues «Evidencias DevSecOps» y «Evidencias AWS» del "
        "repositorio."
    ),
)

titulo("Referencias de esta sección [se integran a la lista general al ensamblar]", 2)
refs = [
    "Amazon Web Services. (s. f.-h). Using the Lambda runtime API for custom runtimes. AWS Lambda Developer Guide. Recuperado el 30 de octubre de 2026, de https://docs.aws.amazon.com/lambda/latest/dg/runtimes-api.html",
    "Aqua Security. (s. f.). Trivy documentation. Recuperado el 30 de octubre de 2026, de https://trivy.dev/",
    "GitHub. (s. f.). About code scanning with CodeQL. GitHub Docs. Recuperado el 30 de octubre de 2026, de https://docs.github.com/en/code-security/code-scanning/introduction-to-code-scanning/about-code-scanning-with-codeql",
    "GitHub. (2026). GHSA-mrhx-c7xc-m3w5: Zod can exhaust memory while validating oversized constrained collections. GitHub Advisory Database. https://github.com/advisories/GHSA-mrhx-c7xc-m3w5",
    "OWASP Mobile Application Security Project. (2024). OWASP Mobile Application Security Verification Standard (MASVS) (versión 2.1.0). OWASP Foundation. https://mas.owasp.org/MASVS/",
    "Semgrep. (s. f.). Testing rules. Semgrep Docs. Recuperado el 30 de octubre de 2026, de https://semgrep.dev/docs/writing-rules/testing-rules",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    for r in p.runs:
        r.font.size = Pt(12)

doc.save(str(SALIDA))
print("Generado:", SALIDA)
